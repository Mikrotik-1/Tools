from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "reports"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "تقرير_نظام_تخزين_الكاميرات.docx"
LOGO = OUT / "it_department_logo.png"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
GREEN = "E9F7EF"
GOLD = "FFF4D6"
RED = "FDECEC"
WHITE = "FFFFFF"
BLACK = "1F2937"


def make_logo():
    image = Image.new("RGBA", (420, 150), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((8, 8, 142, 142), radius=28, fill="#17365D")
    draw.rounded_rectangle((35, 48, 108, 100), radius=8, fill="white")
    draw.polygon([(108, 59), (133, 48), (133, 100), (108, 89)], fill="#7DB7DD")
    draw.ellipse((62, 62, 88, 88), fill="#2E74B5")
    draw.text((170, 42), "IT", fill="#17365D", font=None)
    draw.text((170, 78), "DEPARTMENT", fill="#667085", font=None)
    image.save(LOGO)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, size=11, bold=False, color=BLACK, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    return paragraph


def add_text(doc, text, size=11, bold=False, color=BLACK, after=6, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = rtl(doc.add_paragraph(), align)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    rtl(p)
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else NAVY)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    rtl(p)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.right_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_run(p.add_run(text), size=11)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = rtl(cell.paragraphs[0])
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(label + "  "), size=11, bold=True, color=accent)
    set_run(p.add_run(text), size=11, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def fill_table(table, headers, rows, widths, row_fills=None):
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        p = rtl(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(header), size=10.5, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for r_idx, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for c_idx, value in enumerate(values):
            if row_fills and r_idx - 1 < len(row_fills):
                set_cell_shading(cells[c_idx], row_fills[r_idx - 1])
            p = rtl(cells[c_idx].paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
            set_run(p.add_run(str(value)), size=10)
    set_table_geometry(table, widths)


make_logo()
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for level, size, before, after in ((1, 16, 12, 6), (2, 13, 10, 5), (3, 12, 8, 4)):
    style = styles[f"Heading {level}"]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else NAVY)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Header
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.add_run().add_picture(str(LOGO), width=Inches(1.65))

# Footer with page field
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(fp.add_run("تقرير نظام تخزين كاميرات المراقبة  |  صفحة "), size=8.5, color=MID_GRAY)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
fp._p.append(field)

add_text(doc, "تقرير فني وإداري", size=11, bold=True, color=BLUE, after=2)
title = add_text(doc, "تقييم مدة تخزين تسجيلات كاميرات المراقبة", size=23, bold=True, color=NAVY, after=5)
subtitle = add_text(doc, "تحليل أثر إعادة توزيع الكاميرات والتوصية بزيادة السعة التخزينية", size=12.5, color=MID_GRAY, after=13)

meta = doc.add_table(rows=4, cols=2)
meta_rows = [
    ("موجّه إلى", "السادة رؤساء الإدارات"),
    ("إعداد", "مدير تقنية المعلومات (IT Manager)"),
    ("التاريخ", "24 سبتمبر 2026"),
    ("درجة المستند", "داخلي - لاتخاذ قرار شراء"),
]
for i, (label, value) in enumerate(meta_rows):
    set_cell_shading(meta.cell(i, 0), LIGHT_GRAY)
    set_cell_shading(meta.cell(i, 1), WHITE)
    p0 = rtl(meta.cell(i, 0).paragraphs[0])
    p1 = rtl(meta.cell(i, 1).paragraphs[0])
    set_run(p0.add_run(label), size=10, bold=True, color=NAVY)
    set_run(p1.add_run(value), size=10, color=BLACK)
set_table_geometry(meta, [2000, 7360])

add_heading(doc, "الملخص التنفيذي", 1)
add_text(doc, "بعد إعادة توزيع الكاميرات بين جهازي التسجيل، تحسنت مدة الاحتفاظ بالتسجيلات على الجهاز سعة 64 قناة من 9 أيام إلى 14 يومًا، بينما انخفضت المدة على الجهاز سعة 32 قناة من نحو شهرين ونصف إلى نحو شهر ونصف. وتشير النتيجة إلى انتقال جزء مؤثر من حمل التسجيل إلى جهاز 32 قناة، بما أدى إلى عدم توازن مدة الاحتفاظ بين الجهازين.", after=7)
add_callout(doc, "القرار المطلوب:", "اعتماد شراء قرص تخزين مخصص لأنظمة المراقبة بسعة 10 تيرابايت، بعد التأكد من توافقه مع جهاز التسجيل، لتحسين مدة الاحتفاظ وتقليل الحاجة إلى توسعة جديدة على المدى القريب.", fill=GOLD, accent="7A5A00")

add_heading(doc, "الوضع قبل وبعد نقل الكاميرات", 1)
comparison = doc.add_table(rows=1, cols=5)
comparison_rows = [
    ("جهاز 64 قناة", "9 أيام", "14 يومًا", "+5 أيام", "تحسن 55.6%"),
    ("جهاز 32 قناة", "نحو 75 يومًا", "نحو 45 يومًا", "-30 يومًا", "انخفاض 40%"),
]
fill_table(comparison, ["جهاز التسجيل", "قبل النقل", "بعد النقل", "التغير", "الدلالة"], comparison_rows, [1960, 1500, 1500, 1500, 2900], [GREEN, RED])

add_text(doc, "ملاحظة حسابية: تم تحويل شهرين ونصف إلى نحو 75 يومًا، وشهر ونصف إلى نحو 45 يومًا لأغراض المقارنة الإدارية.", size=8.5, color=MID_GRAY, after=4)

add_heading(doc, "قراءة النتائج", 1)
add_bullet(doc, "جهاز 64 قناة أصبح يحتفظ بالتسجيلات لمدة أطول بخمسة أيام، وهو تحسن واضح ناتج عن انخفاض حمل التسجيل عليه بعد نقل بعض الكاميرات.")
add_bullet(doc, "جهاز 32 قناة فقد نحو 30 يومًا من مدة الاحتفاظ، ما يجعله نقطة الضغط الحالية في منظومة التخزين.")
add_bullet(doc, "الوضع الحالي يسمح بالتسجيل، لكنه يقلل الفترة المتاحة للرجوع إلى الأحداث على جهاز 32 قناة مقارنة بالوضع السابق.")
add_bullet(doc, "مدة التسجيل الفعلية تتأثر بعدد الكاميرات ودقتها ومعدل الإطارات والـBitrate ونوع الضغط ونمط التسجيل المستمر أو بالحركة.")

add_heading(doc, "بدائل زيادة السعة والتكلفة التقديرية", 1)
options = doc.add_table(rows=1, cols=5)
option_rows = [
    ("8 تيرابايت", "20,600 جنيه", "2,575 جنيه/تيرا", "تكلفة أولية أقل", "مقبول"),
    ("10 تيرابايت", "22,000 جنيه", "2,200 جنيه/تيرا", "سعة أعلى 25%", "موصى به"),
]
fill_table(options, ["السعة", "السعر التقريبي", "تكلفة التيرابايت", "الميزة", "التقييم"], option_rows, [1600, 1900, 1900, 2200, 1760], [WHITE, GREEN])
add_text(doc, "فرق السعر بين الخيارين نحو 1,400 جنيه فقط (قرابة 6.8% من سعر 8 تيرا)، بينما يوفر خيار 10 تيرا سعة إضافية قدرها 25% وتكلفة أقل لكل تيرابايت بنحو 14.6%.", size=10, color=MID_GRAY, after=5)

add_heading(doc, "التوصية الفنية", 1)
add_callout(doc, "التوصية:", "شراء هارد 10 تيرابايت من فئة Surveillance ومصمم للعمل المتواصل 24/7. هذا الخيار يحقق أفضل قيمة مالية ويمنح هامشًا أكبر لاستيعاب أي زيادة مستقبلية في عدد الكاميرات أو جودة التسجيل.", fill=LIGHT_BLUE, accent=NAVY)
add_bullet(doc, "التأكد قبل الشراء من الحد الأقصى للسعة المدعومة في جهاز التسجيل ومن إصدار الـFirmware.")
add_bullet(doc, "اختيار قرص أصلي بضمان معتمد وتقنية CMR، وتجنب الأقراص غير المخصصة للتسجيل المستمر.")
add_bullet(doc, "تركيب القرص وتهيئته من داخل جهاز التسجيل، ثم متابعة مدة الاحتفاظ الفعلية لمدة أسبوعين على الأقل.")
add_bullet(doc, "مراجعة إعدادات الدقة وBitrate والضغط H.265 ونمط التسجيل، دون خفض الجودة بما يؤثر على الاستفادة الأمنية من التسجيلات.")

add_heading(doc, "خطة التنفيذ المقترحة", 1)
steps = doc.add_table(rows=1, cols=4)
step_rows = [
    ("1", "مراجعة توافق السعة ونوع القرص", "إدارة تقنية المعلومات", "قبل الشراء"),
    ("2", "الحصول على عرض سعر وضمان رسمي", "المشتريات / IT", "يوم عمل"),
    ("3", "الشراء والتركيب والتهيئة", "إدارة تقنية المعلومات", "حسب التوريد"),
    ("4", "قياس مدة التسجيل بعد التركيب", "إدارة تقنية المعلومات", "بعد 14 يومًا"),
]
fill_table(steps, ["المرحلة", "الإجراء", "المسؤول", "التوقيت"], step_rows, [1000, 3900, 2460, 2000])

add_heading(doc, "الاعتماد", 1)
add_text(doc, "نرجو التكرم بالموافقة على شراء هارد سعة 10 تيرابايت، على أن يتم التنفيذ بعد التحقق الفني النهائي من التوافق والحصول على عرض سعر رسمي.", after=14)

signature = doc.add_table(rows=4, cols=2)
signature_data = [
    ("إعداد التقرير", "اعتماد الإدارة"),
    ("الاسم: ........................................", "الاسم: ........................................"),
    ("الصفة: مدير تقنية المعلومات (IT Manager)", "الصفة: ........................................"),
    ("التوقيع والتاريخ: ............................", "التوقيع والتاريخ: ............................"),
]
for r_idx, row in enumerate(signature_data):
    for c_idx, text in enumerate(row):
        cell = signature.cell(r_idx, c_idx)
        if r_idx == 0:
            set_cell_shading(cell, NAVY)
        p = rtl(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
        set_run(p.add_run(text), size=10, bold=(r_idx == 0), color=WHITE if r_idx == 0 else BLACK)
set_table_geometry(signature, [4680, 4680])

doc.core_properties.title = "تقرير تقييم نظام تخزين كاميرات المراقبة"
doc.core_properties.subject = "تحليل مدة الاحتفاظ قبل وبعد إعادة توزيع الكاميرات"
doc.core_properties.author = "إدارة تقنية المعلومات"
doc.core_properties.keywords = "كاميرات مراقبة، تخزين، NVR، DVR، هارد Surveillance"
doc.save(DOCX)
print(DOCX)
